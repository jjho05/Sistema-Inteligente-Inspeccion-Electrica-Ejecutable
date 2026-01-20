"""
Document generator for creating official Word inspection reports.
Uses the official template and fills it with analysis data.
"""

from docx import Document
from pathlib import Path
from typing import Dict, Any
from datetime import datetime

from backend.utils.config import TEMPLATES_PATH, GENERATED_PATH


class DocumentGenerator:
    """Generates Word documents from inspection data."""
    
    def __init__(self, template_path: str = None):
        """
        Initialize document generator.
        
        Args:
            template_path: Path to Word template (optional)
        """
        if template_path is None:
            # Use cleaned template (without personal contact info)
            template_path = Path(TEMPLATES_PATH) / "DictamenElectrico_Cleaned.docx"
        
        self.template_path = Path(template_path)
        if not self.template_path.exists():
            raise FileNotFoundError(f"Template not found: {template_path}")
        
        self.generated_path = Path(GENERATED_PATH)
        self.generated_path.mkdir(parents=True, exist_ok=True)
    
    def generate_dictamen(self, dictamen_data: Dict[str, Any], 
                         output_filename: str = None) -> str:
        """
        Generate dictamen Word document.
        
        Args:
            dictamen_data: Data to fill in the template
            output_filename: Optional output filename
            
        Returns:
            Path to generated document
        """
        print("Generating Word document...")
        
        # Load template
        doc = Document(self.template_path)
        
        # Fill template
        self._fill_template(doc, dictamen_data)
        
        # Generate output filename if not provided
        if output_filename is None:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            folio = dictamen_data.get('folio', 'SIN_FOLIO')
            output_filename = f"Dictamen_{folio}_{timestamp}.docx"
        
        # Save document
        output_path = self.generated_path / output_filename
        doc.save(str(output_path))
        
        print(f"✓ Document generated: {output_path}")
        
        return str(output_path)
    
    def _fill_template(self, doc: Document, data: Dict[str, Any]):
        """Fill template with data using template mapper."""
        from backend.utils.template_mapper import map_data_to_template
        
        try:
            # Use the template mapper for proper field mapping
            map_data_to_template(doc, data)
        except Exception as e:
            print(f"Warning: Error using template mapper: {e}")
            # Fallback to simple placeholder replacement
            self._simple_fill(doc, data)
    
    def _simple_fill(self, doc: Document, data: Dict[str, Any]):
        """Simple fallback method for filling template."""
        # Replace placeholders in paragraphs
        for paragraph in doc.paragraphs:
            self._replace_placeholder(paragraph, '{{FOLIO}}', data.get('folio', ''))
            self._replace_placeholder(paragraph, '{{FECHA}}', data.get('fecha', ''))
            self._replace_placeholder(paragraph, '{{SOLICITANTE}}', data.get('solicitante', ''))
            self._replace_placeholder(paragraph, '{{UBICACION}}', data.get('ubicacion', ''))
            self._replace_placeholder(paragraph, '{{TIPO}}', data.get('tipo_instalacion', ''))
            self._replace_placeholder(paragraph, '{{OBSERVACIONES}}', data.get('observaciones', ''))
            self._replace_placeholder(paragraph, '{{DICTAMEN}}', data.get('dictamen_final', ''))
    
    def _replace_placeholder(self, paragraph, placeholder: str, value: str):
        """Replace placeholder in paragraph."""
        if placeholder in paragraph.text:
            paragraph.text = paragraph.text.replace(placeholder, value)


def generate_inspection_document(dictamen_data: Dict[str, Any], 
                                 output_filename: str = None) -> str:
    """
    Convenience function to generate inspection document.
    
    Args:
        dictamen_data: Dictamen data
        output_filename: Optional output filename
        
    Returns:
        Path to generated document
    """
    generator = DocumentGenerator()
    return generator.generate_dictamen(dictamen_data, output_filename)
